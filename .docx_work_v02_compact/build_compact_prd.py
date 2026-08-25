from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\wenbo\OneDrive\文档\ChatGPT\飞猪ToC_AI助手")
SOURCE = ROOT / "飞猪AI旅行规划助手_PRD_V0.2.docx"
OUTPUT = ROOT / "飞猪AI旅行规划助手_PRD_V0.2_精简版.docx"

BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
TEXT = "22303C"
MUTED = "667085"
WHITE = "FFFFFF"
USABLE_DXA = 9360


def set_run_font(run, name="微软雅黑", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def format_cell_text(cell, bold=False, color=TEXT, size=9.5, align=None):
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        if align is not None:
            p.alignment = align
        for run in p.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths_dxa, first_col_label=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
        format_cell_text(table.rows[0].cells[i], bold=True, color=TEXT)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            row.cells[i].text = value
            if first_col_label and i == 0:
                set_cell_shading(row.cells[i], LIGHT_GRAY)
                format_cell_text(row.cells[i], bold=True, color=TEXT)
            else:
                format_cell_text(row.cells[i])
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        lead, rest = text[: len(bold_lead)], text[len(bold_lead) :]
        r = p.add_run(lead)
        set_run_font(r, size=10.5, bold=True, color=TEXT)
        r = p.add_run(rest)
        set_run_font(r, size=10.5, color=TEXT)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=TEXT)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_together = True
        r = p.add_run(item)
        set_run_font(r, size=10.2, color=TEXT)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_together = True
        r = p.add_run(item)
        set_run_font(r, size=10.2, color=TEXT)


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title + "  ")
    set_run_font(r, size=10, bold=True, color=BLUE)
    r = p.add_run(text)
    set_run_font(r, size=10, color=TEXT)
    set_table_geometry(table, [USABLE_DXA])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def set_dynamic_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("飞猪 AI 旅行规划助手｜PRD V0.2 精简版")
    set_run_font(r, size=9, color=MUTED)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 4)):
        style = doc.styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(10.2)
        style.font.color.rgb = RGBColor.from_string(TEXT)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("飞猪 AI 旅行规划助手｜产品需求文档")
    set_run_font(hr, size=9, color=MUTED)

    doc.core_properties.title = "飞猪 AI 旅行规划助手 PRD V0.2（精简版）"
    doc.core_properties.subject = "To C 手机端单页面 AI 旅行规划 Demo"
    doc.core_properties.comments = "V0.2 compact edition"

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    set_run_font(r, size=10, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("飞猪 AI 旅行规划助手")
    set_run_font(r, size=28, bold=True, color=TEXT)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("To C 手机端单页面 AI 旅行规划 Demo")
    set_run_font(r, size=14, color=MUTED)

    add_table(
        doc,
        ["项目", "内容"],
        [
            ("文档编号", "PRD-FLIGGY-AI-001"),
            ("文档版本", "V0.2（精简版）"),
            ("文档状态", "待确认"),
            ("编写部门", "产品部"),
            ("更新日期", "2026-08-19"),
            ("适用范围", "手机端单页面 MVP / Demo"),
        ],
        [2200, 7160],
        first_col_label=True,
    )
    add_callout(
        doc,
        "核心范围",
        "只开发一个移动端 AI 对话页面，完成需求理解、主动追问、结构化行程生成和最小必要修改。",
    )
    add_body(doc, "文档用途：用于产品范围确认、原型设计、Demo 开发与功能验收。")

    add_heading(doc, "修订记录", 1)
    add_table(
        doc,
        ["版本", "日期", "主要内容"],
        [
            ("V0.1", "2026-08-19", "定义单页面 AI 旅行规划 MVP。"),
            ("V0.2", "2026-08-19", "支持单双目的地；完善追问、修改、价格与 Demo 验收规则；合并重复描述。"),
        ],
        [1300, 1900, 6160],
    )

    # 1
    add_heading(doc, "1. 产品定义与目标", 1)
    add_body(
        doc,
        "产品定位：飞猪 AI 旅行规划助手是一款 To C 手机端 AI Native Demo。用户通过自然语言描述需求，AI 在同一聊天页面内完成信息收集、行程生成和后续修改。",
        "产品定位：",
    )
    add_bullets(
        doc,
        [
            "支持单目的地和多目的地规划；多目的地路线为重点展示能力，长沙三日游为回归场景。",
            "将模糊旅行想法转化为清晰、结构化、可继续修改的每日行程。",
            "通过少量主动追问补齐关键信息，避免把对话变成表单填写。",
            "第一版以稳定 Demo 为目标，使用 Mock 数据验证完整产品闭环。",
        ],
    )

    # 2
    add_heading(doc, "2. MVP 范围", 1)
    add_table(
        doc,
        ["范围内", "范围外"],
        [
            ("一个手机端 AI 对话页面", "登录、注册、首页、个人中心"),
            ("自然语言需求识别与多轮追问", "订单、支付、真实交易闭环"),
            ("单/多目的地结构化行程生成", "酒店、机票、景点独立列表或详情页"),
            ("按天或按全局条件修改行程", "地图、社区、收藏、多页面路由"),
            ("Mock 数据及参考价格展示", "未接入数据源的实时价格、库存和优惠"),
        ],
        [4680, 4680],
    )
    add_callout(doc, "范围原则", "第一版不展示尚未实现的外围功能入口，不承诺实时数据。")

    # 3
    add_heading(doc, "3. 核心业务流程与状态", 1)
    add_body(
        doc,
        "用户进入 → AI 欢迎 → 用户描述需求 → AI 抽取并保存字段 → 判断完整性 → 主动追问 → 生成结构化行程 → 用户提出修改 → AI 按影响范围更新。",
    )
    add_table(
        doc,
        ["状态", "UI / AI 行为", "进入条件"],
        [
            ("INITIAL", "欢迎语与快捷建议", "首次进入"),
            ("COLLECTING_REQUIREMENTS", "解析用户输入并累计字段", "收到需求或补充信息"),
            ("ASKING", "每轮追问 1–2 个问题", "关键字段缺失或冲突"),
            ("READY_TO_GENERATE", "确认条件满足", "关键字段完整"),
            ("GENERATING", "显示“正在规划你的旅行…”", "开始生成"),
            ("PLAN_GENERATED", "展示结构化行程卡片", "生成成功"),
            ("MODIFYING", "识别影响范围并更新", "已有行程且收到修改"),
            ("ERROR", "说明原因，保留有效内容并允许重试", "生成或修改失败"),
        ],
        [2500, 4300, 2560],
    )

    # 4
    add_heading(doc, "4. 需求收集规则", 1)
    add_table(
        doc,
        ["级别", "字段", "处理规则"],
        [
            ("核心必填", "destination(s)、startDate 或 duration、travelers", "缺失时主动追问，不直接生成行程。"),
            ("条件必填", "departureCity", "多目的地或需要规划跨城交通时询问。"),
            ("推荐补充", "budget、preferences、pace", "可适当追问；用户不提供时使用默认假设并说明。"),
            ("可选", "companion、transport、hotel、constraints", "用于提高推荐质量，不阻塞生成。"),
        ],
        [1800, 3300, 4260],
    )
    add_bullets(
        doc,
        [
            "已经识别的字段必须保存在当前会话中，后续补充时不得清空。",
            "每轮最多追问 1–2 个问题，不重复询问已经获得的信息。",
            "目的地范围过大时，先确认具体城市或路线；用户选择单目的地时不强制增加第二站。",
            "日期、天数、预算或目的地发生冲突时，先指出冲突并等待确认。",
        ],
    )

    # 5
    add_heading(doc, "5. 行程生成与修改规则", 1)
    add_heading(doc, "5.1 行程生成", 2)
    add_bullets(
        doc,
        [
            "使用 TripPlan → DayPlan → Activity 三级结构保存行程。",
            "单目的地按区域和主题集中安排；多目的地同时给出路线顺序、停留天数和城际交通。",
            "每天按时间顺序展示，默认安排 2–3 个主要活动，并根据 pace 控制强度。",
            "结合偏好与预算给出建议；交通、住宿和活动费用均需注明数据属性。",
        ],
    )
    add_heading(doc, "5.2 最小必要修改", 2)
    add_bullets(
        doc,
        [
            "修改某一天：只更新目标日期及其必要的交通和费用，其他日期保持不变。",
            "增删或替换目的地：重新检查受影响的路线、日期、交通和预算。",
            "修改预算、节奏等全局条件：检查全部行程，但只调整不符合条件的部分。",
            "修改后明确说明改动内容；失败时保留最近一次有效行程。",
        ],
    )

    # 6
    add_heading(doc, "6. 页面与交互要求", 1)
    add_table(
        doc,
        ["区域", "内容"],
        [
            ("顶部导航", "返回、标题“AI旅行助手”、AI 在线状态。"),
            ("消息区域", "AI/用户消息、思考状态、主动追问、Chip、行程卡片和更新卡片。"),
            ("底部输入", "固定输入框与发送按钮；Placeholder 为“告诉 AI 你的旅行需求…”。"),
        ],
        [2200, 7160],
        first_col_label=True,
    )
    add_bullets(
        doc,
        [
            "首次进入显示欢迎语，并提供长沙三日游、周末推荐和美食旅行快捷建议。",
            "聊天区域可滚动；新消息、追问和行程卡片出现后自动滚动到底部。",
            "输入为空时不可发送；防止重复发送；生成过程中允许停止。",
            "清空会话必须二次确认；失败后允许重试且不丢失最近一次有效行程。",
        ],
    )

    # 7
    add_heading(doc, "7. 价格数据与 Demo 假设", 1)
    add_body(
        doc,
        "价格规则：AI 不直接生成具有实时属性的价格。景点、住宿、交通和餐饮费用原则上由价格数据服务提供；无法获得实时数据时，使用参考价格或费用区间，并标注“预估价格”或“参考价格”。",
        "价格规则：",
    )
    add_bullets(
        doc,
        [
            "实时机票、酒店等价格仅在接入对应数据源后展示，并提示价格可能随库存和时间变化。",
            "V0.2 Demo 默认假设：儿童不超过 12 岁半价；老人达到 65 岁免费；酒店每间最多入住 2 人。",
            "上述假设仅用于功能验证，不代表所有真实景区、酒店或交通平台政策。",
            "商业产品上线后，实时数据源和供应商实际规则优先于 Demo 默认规则。",
        ],
    )

    # 8
    add_heading(doc, "8. Demo 场景与验收标准", 1)
    add_heading(doc, "8.1 必须跑通的场景", 2)
    add_numbers(
        doc,
        [
            "单目的地：用户输入“我想去长沙玩 3 天” → AI 追问日期和人数 → 补充预算与偏好 → 生成长沙 3 天 2 晚行程 → 用户要求“第二天不要爬山” → 仅更新 Day 2。",
            "多目的地：长沙出发前往张家界、芙蓉镇、凤凰古城 → AI 补齐人数、预算和偏好 → 生成路线顺序、停留天数、城际交通和每日行程 → 支持增删目的地或调整停留时间。",
        ],
    )
    add_heading(doc, "8.2 MVP 验收", 2)
    add_bullets(
        doc,
        [
            "能识别单个或多个目的地，并持续保存已获取的会话字段。",
            "缺少核心字段时不直接生成，每轮最多追问 1–2 项。",
            "生成结果包含摘要、按天活动、交通提示、强度和预估费用。",
            "局部修改不影响无关日期；全局修改只调整违反新条件的安排。",
            "所有非实时费用和优惠政策均有清晰标识。",
            "生成或修改失败后可恢复，且最近一次有效行程仍可使用。",
            "全部核心流程在同一页面完成，390 × 844 px 下无横向溢出。",
        ],
    )

    # 9
    add_heading(doc, "9. 非功能约束与文档边界", 1)
    add_bullets(
        doc,
        [
            "Mock Demo：发送后 0.5 秒内出现反馈，完整结果目标 3 秒内返回。",
            "优先保证手机端 Chrome、Safari 和常见 WebView 的基本可用性。",
            "UI 与 AI/对话逻辑分离，前端不得硬编码 API Key。",
            "PRD 只描述产品范围、用户体验、业务规则和验收标准。",
            "项目目录、组件拆分、Service 设计、Schema、API 契约、测试和部署方案另行编写《开发技术方案》。",
        ],
    )
    add_callout(doc, "版本结论", "V0.2 以核心 Demo 闭环为唯一优先级，不新增登录、交易、地图或多页面业务模块。")

    set_dynamic_footer(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
